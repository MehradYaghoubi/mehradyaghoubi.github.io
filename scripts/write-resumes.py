"""Build readable DOCX resumes with explicit paragraph and run direction."""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SPECS = json.loads((ROOT / 'tmp/resumes/resume-data.json').read_text(encoding='utf-8'))
OUT = ROOT / 'public/resumes'
PERSIAN = re.compile(r'[\u0621-\u064a\u067e\u0686\u0698\u06a9\u06af\u06cc]')
# Include Latin parentheses in the same run as their text, so they cannot reverse.
LATIN = re.compile(r'[A-Za-z0-9][A-Za-z0-9._:/@#?&=+%-]*(?:[ \t]+(?:\([A-Za-z0-9][A-Za-z0-9 ._:/@#?&=+%-]*\)|[A-Za-z0-9][A-Za-z0-9._:/@#?&=+%-]*))*')


def flag(parent, key, value):
    element = OxmlElement('w:' + key)
    element.set(qn('w:val'), str(value))
    parent.append(element)


def run_properties(parent, font, size, bold, rtl, color='17212B'):
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
        fonts.set(qn('w:' + key), font)
    props.append(fonts)
    for key, value in [('sz', int(size * 2)), ('szCs', int(size * 2)),
                       ('b', int(bold)), ('bCs', int(bold)), ('rtl', int(rtl)), ('color', color)]:
        flag(props, key, value)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'fa-IR' if rtl else 'en-US')
    lang.set(qn('w:bidi'), 'fa-IR' if rtl else 'en-US')
    props.append(lang)
    parent.insert(0, props)


def add_text(p, text, fa, size, bold=False, color='17212B'):
    pieces = []
    if fa:
        start = 0
        for match in LATIN.finditer(text):
            pieces.extend([(text[start:match.start()], True), ('\u200e' + match.group() + '\u200e', False)])
            start = match.end()
        pieces.append((text[start:], True))
    else:
        pieces = [(text, False)]
    for value, rtl in pieces:
        if not value:
            continue
        r = p.add_run(value)
        run_properties(r._r, 'Vazirmatn' if rtl else 'Arial', size, bold, rtl, color)


def para(d, text='', *, fa=False, latin=False, size=None, bold=False, after=4, before=0,
         keep=False, style=None, color='17212B'):
    p = d.add_paragraph(style=style)
    flag(p._p.get_or_add_pPr(), 'bidi', int(fa))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if fa else WD_ALIGN_PARAGRAPH.LEFT
    f = p.paragraph_format
    font_size = size or (11.5 if fa else 11)
    f.line_spacing = Pt(font_size * (1.5 if fa and not latin else 1.35))
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.keep_with_next, f.keep_together, f.widow_control = keep, True, True
    add_text(p, text, fa and not latin, size or (11.5 if fa else 11), bold, color)
    return p


def links(d, items, fa, size=10.5, after=7):
    p = para(d, fa=fa, after=after)
    for i, (label, target) in enumerate(items):
        if i:
            add_text(p, '  ·  ', fa, size)
        rtl = bool(PERSIAN.search(label))
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        run_properties(run, 'Vazirmatn' if rtl else 'Arial', size, False, rtl, '175A82')
        text = OxmlElement('w:t')
        text.text = label if rtl or not fa else '\u200e' + label + '\u200e'
        run.append(text)
        link.append(run)
        p._p.append(link)
    return p


for spec in SPECS:
    fa = spec['lang'] == 'fa'
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(16)
    s.left_margin = s.right_margin = Mm(18)
    s.footer_distance = Mm(8)
    normal = d.styles['Normal']
    normal.font.name = 'Vazirmatn' if fa else 'Arial'
    normal.font.size = Pt(11.5 if fa else 11)
    for name in ['Title', 'Heading 1']:
        d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in d.styles[name].element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    flag(d.settings.element, 'embedTrueTypeFonts', 1)
    flag(d.settings.element, 'saveSubsetFonts', 0)
    para(d, spec['name'], fa=fa, size=25, bold=True, after=3, style='Title', color='000000', keep=True)
    para(d, spec['title'], fa=fa, size=13, bold=True, after=4, color='000000', keep=True)
    if fa:
        para(d, 'Mehrad Yaghoubi', fa=True, latin=True, size=10.5, after=3, keep=True)
    para(d, spec['location'], fa=fa, size=10.5, after=3, keep=True)
    links(d, [(spec['email'], 'mailto:' + spec['email']), (spec['phone'], 'tel:+989912093038')], fa, after=4)
    links(d, spec['links'], fa, after=7)

    for section in spec['sections']:
        heading = para(d, section['heading'], fa=fa, size=13.5, bold=True, before=9, after=5,
                       keep=True, style='Heading 1', color='000000')
        heading.paragraph_format.page_break_before = section.get('pageBreakBefore', False)
        for entry in section['entries']:
            if 'text' in entry:
                para(d, entry['text'], fa=fa, latin=entry.get('latin', False), after=5)
                continue
            start = len(d.paragraphs)
            if entry.get('title'):
                title = entry['title']
                para(d, title, fa=fa, latin=fa and not PERSIAN.search(title), bold=True, after=2)
            if entry.get('meta'):
                para(d, entry['meta'], fa=fa, size=10.5, after=3, color='334155')
            for value in entry.get('paragraphs', []):
                para(d, value, fa=fa, after=4)
            for value in entry.get('bullets', []):
                para(d, '• ' + value, fa=fa, after=4)
            if entry.get('status'):
                para(d, entry['status'], fa=fa, size=10.5, after=2, color='334155')
            if entry.get('statusLatin'):
                para(d, entry['statusLatin'], fa=fa, latin=True, size=10.5, after=3, color='334155')
            if entry.get('links'):
                links(d, entry['links'], fa)
            paragraphs = d.paragraphs[start:]
            for p in paragraphs[:-1]:
                p.paragraph_format.keep_with_next = True
            if paragraphs:
                paragraphs[-1].paragraph_format.space_after = Pt(7)

    footer = s.footer.paragraphs[0]
    flag(footer._p.get_or_add_pPr(), 'bidi', 0)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, 'Mehrad Yaghoubi  ·  ', False, 9)
    for i, name in enumerate(['PAGE', 'NUMPAGES']):
        if i:
            add_text(footer, ' / ', False, 9)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), name)
        footer._p.append(field)
    d.core_properties.title = spec['name'] + ' ' + spec['title']
    d.core_properties.author = 'Mehrad Yaghoubi'
    d.save(OUT / (spec['filename'] + '.docx'))
    print(spec['filename'] + '.docx')
